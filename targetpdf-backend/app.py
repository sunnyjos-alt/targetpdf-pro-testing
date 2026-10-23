from flask import Flask, request, send_file, jsonify
from flask_cors import CORS
from pdf2docx import Converter
from docx import Document
from docx.shared import Inches, Pt
from io import BytesIO
from pathlib import Path
from PIL import Image
import fitz
import logging
import os
import shutil
import tempfile
import time

try:
    import pytesseract
    OCR_AVAILABLE = True
except Exception:
    pytesseract = None
    OCR_AVAILABLE = False

app = Flask(__name__)
app.logger.setLevel(logging.INFO)

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_BYTES

ALLOWED_ORIGINS = [
    "https://targetpdf-pro-testing.pages.dev",
    "https://targetpdf.com",
    "https://www.targetpdf.com",
]
CORS(
    app,
    resources={r"/*": {"origins": ALLOWED_ORIGINS}},
    expose_headers=[
        "X-Conversion-Seconds", "X-Input-Bytes", "Content-Disposition",
        "X-Conversion-Mode", "X-Conversion-Notes", "X-PDF-Pages",
    ],
)

@app.get("/")
@app.get("/health")
def health():
    return jsonify({
        "status": "ok",
        "service": "TargetPDF PDF-to-Word backend",
        "ocr_available": bool(OCR_AVAILABLE and shutil.which("tesseract")),
        "pytesseract_imported": bool(OCR_AVAILABLE),
        "tesseract_executable": shutil.which("tesseract"),
        "features": ["text-pdf-conversion", "pdf-page-analysis", "scanned-pdf-ocr"],
    }), 200

@app.errorhandler(413)
def request_entity_too_large(_error):
    return jsonify({
        "error": "The PDF is too large. Please upload a PDF smaller than 20 MB."
    }), 413


def inspect_pdf(pdf_path):
    """Measure selectable text page-by-page without rendering page images."""
    pdf = fitz.open(pdf_path)
    page_texts = []
    for page in pdf:
        try:
            page_texts.append(page.get_text("text", sort=True).strip())
        except TypeError:
            page_texts.append(page.get_text("text").strip())
    page_count = len(pdf)
    pdf.close()
    if not page_count:
        raise ValueError("This PDF contains no pages.")
    text_pages = sum(1 for text in page_texts if len(text) >= 40)
    scan_pages = page_count - text_pages
    return page_texts, page_count, text_pages, scan_pages


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)


def add_page_text(doc, text):
    # Keep paragraph boundaries while avoiding hundreds of empty paragraphs.
    for block in text.split("\n\n"):
        cleaned = "\n".join(line.rstrip() for line in block.splitlines()).strip()
        if cleaned:
            doc.add_paragraph(cleaned)


def page_text_or_ocr(page, extracted_text):
    # Keep selectable text when it appears substantial; OCR image-only or nearly
    # image-only pages. This avoids duplicate text on common mixed pages.
    if len(extracted_text.strip()) >= 40:
        return extracted_text.strip(), False
    if not (OCR_AVAILABLE and shutil.which("tesseract")):
        return extracted_text.strip(), False
    pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), colorspace=fitz.csGRAY, alpha=False)
    image = Image.open(BytesIO(pix.tobytes("png")))
    text = pytesseract.image_to_string(image, lang="eng", config="--psm 3")
    return text.strip(), True


def convert_with_ocr_or_text(pdf_path, page_texts, page_count):
    """Fallback for scanned/mixed PDFs: create editable text in original page order."""
    if not (OCR_AVAILABLE and shutil.which("tesseract")):
        raise RuntimeError(
            "This PDF appears to contain scanned pages, but OCR is not available on the server yet."
        )
    doc = Document()
    configure_document(doc)
    pdf = fitz.open(pdf_path)
    ocr_pages = 0
    pages_with_text = 0
    try:
        for index, page in enumerate(pdf):
            extracted = page_texts[index] if index < len(page_texts) else ""
            page_text, used_ocr = page_text_or_ocr(page, extracted)
            if used_ocr:
                ocr_pages += 1
            if page_text:
                add_page_text(doc, page_text)
                pages_with_text += 1
            else:
                doc.add_paragraph("[No readable text could be recognized on this page.]")
            if index < page_count - 1:
                doc.add_page_break()
    finally:
        pdf.close()

    output = BytesIO()
    doc.save(output)
    data = output.getvalue()
    if not data:
        raise RuntimeError("OCR/text conversion produced an empty Word document.")
    mode = "OCR-assisted editable text"
    notes = (
        f"Processed {page_count} pages; OCR was used on {ocr_pages} pages. "
        "This fallback prioritizes editable text and page order; complex formatting, "
        "columns, tables, and images may need manual adjustment."
    )
    return data, mode, notes, pages_with_text


def convert_with_pdf2docx(pdf_path, docx_path):
    converter = Converter(pdf_path)
    try:
        converter.convert(docx_path, start=0, end=None)
    finally:
        converter.close()
    if not os.path.isfile(docx_path) or os.path.getsize(docx_path) == 0:
        raise RuntimeError("Converter produced an empty DOCX file")
    with open(docx_path, "rb") as docx_file:
        data = docx_file.read()
    # Verify the DOCX can be opened and contains some editable text.
    check = Document(BytesIO(data))
    text_chars = sum(len(p.text.strip()) for p in check.paragraphs)
    for table in check.tables:
        for row in table.rows:
            for cell in row.cells:
                text_chars += len(cell.text.strip())
    return data, text_chars


@app.post("/convert-to-word")
def convert_to_word():
    request_started = time.perf_counter()
    if "file" not in request.files:
        return jsonify({"error": "No file was uploaded."}), 400

    uploaded = request.files["file"]
    if not uploaded.filename:
        return jsonify({"error": "Please choose a PDF file."}), 400
    if Path(uploaded.filename).suffix.lower() != ".pdf":
        return jsonify({"error": "Please upload a PDF file."}), 415

    pdf_bytes = uploaded.read()
    input_bytes = len(pdf_bytes)
    if input_bytes == 0:
        return jsonify({"error": "The selected PDF is empty."}), 400
    if not pdf_bytes.startswith(b"%PDF-"):
        return jsonify({"error": "This file does not appear to be a valid PDF."}), 415

    conversion_started = time.perf_counter()
    try:
        with tempfile.TemporaryDirectory(prefix="targetpdf-") as temp_dir:
            pdf_path = os.path.join(temp_dir, "input.pdf")
            docx_path = os.path.join(temp_dir, "converted.docx")
            with open(pdf_path, "wb") as pdf_file:
                pdf_file.write(pdf_bytes)

            page_texts, page_count, text_pages, scan_pages = inspect_pdf(pdf_path)
            text_ratio = text_pages / page_count

            if text_ratio >= 0.98:
                docx_bytes, output_chars = convert_with_pdf2docx(pdf_path, docx_path)
                mode = "PDF layout conversion"
                notes = (
                    f"Processed {page_count} pages. Selectable text was detected on "
                    f"{text_pages} pages. Review the Word file for layout differences."
                )
                # A DOCX can exist yet contain almost no editable text. Do not call
                # that a successful text conversion without warning the user.
                input_chars = sum(len(t) for t in page_texts)
                if input_chars > 300 and output_chars < max(100, int(input_chars * 0.08)):
                    raise RuntimeError(
                        "The PDF contains text, but very little editable text was found in the Word output. "
                        "The file may use unusual fonts or a layout this converter cannot read reliably."
                    )
            else:
                docx_bytes, mode, notes, pages_with_text = convert_with_ocr_or_text(
                    pdf_path, page_texts, page_count
                )
                notes += f" Editable text was produced for {pages_with_text} of {page_count} pages."
                if text_ratio >= 0.15:
                    notes += " This was a mixed PDF; text pages were extracted and image-only pages were OCR-processed."
                else:
                    notes += " This PDF was mostly scanned; OCR recognition depends on scan clarity."
        
        conversion_seconds = time.perf_counter() - conversion_started
        total_seconds = time.perf_counter() - request_started
        app.logger.info(
            "PDF-to-Word success input_bytes=%d output_bytes=%d pages=%d mode=%s seconds=%.2f",
            input_bytes, len(docx_bytes), page_count, mode, conversion_seconds
        )

        response = send_file(
            BytesIO(docx_bytes),
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            as_attachment=True,
            download_name="converted.docx",
        )
        response.headers["X-Conversion-Seconds"] = f"{conversion_seconds:.2f}"
        response.headers["X-Input-Bytes"] = str(input_bytes)
        response.headers["X-Conversion-Mode"] = mode
        response.headers["X-Conversion-Notes"] = notes
        response.headers["X-PDF-Pages"] = str(page_count)
        response.headers["Cache-Control"] = "no-store"
        return response

    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        app.logger.exception(
            "PDF-to-Word failure input_bytes=%d elapsed_seconds=%.2f",
            input_bytes, time.perf_counter() - request_started
        )
        # Give a useful explanation without exposing internal paths or stack traces.
        message = str(exc).strip()
        if not message or len(message) > 350:
            message = (
                "Conversion failed. The PDF may be damaged, password-protected, "
                "too complex, or too large for the free server's resources."
            )
        return jsonify({"error": message}), 500

if __name__ == "__main__":
    port = int(os.environ.get("PORT", "10000"))
    app.run(host="0.0.0.0", port=port)
